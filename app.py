"""
app.py
Streamlit UI for the AI Job Application Agent.
Lets the user upload/paste a resume + paste a JD, runs the LangGraph pipeline,
and displays skill gap analysis, cover letter, resume bullets, and an
optional Q&A assistant for application questions.
"""

import streamlit as st
from graph.workflow import run_pipeline
from pypdf import PdfReader
from docx import Document
import io

st.set_page_config(page_title="AI Job Application Agent", page_icon="🎯", layout="wide")


def extract_text_from_file(uploaded_file) -> str:
    """Extracts plain text from an uploaded PDF or DOCX file."""
    if uploaded_file is None:
        return ""

    file_bytes = uploaded_file.read()
    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(file_bytes))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        return text.strip()

    elif filename.endswith(".docx"):
        doc = Document(io.BytesIO(file_bytes))
        text = "\n".join(para.text for para in doc.paragraphs)
        return text.strip()

    else:
        st.error("Unsupported file type. Please upload a PDF or DOCX.")
        return ""


st.title("🎯 AI Job Application Agent")
st.caption("Upload or paste your resume and a job description — get a skill-gap analysis, a tailored cover letter, and ATS-optimized resume bullets.")

# ---- Session state so results survive re-renders ----
if "result" not in st.session_state:
    st.session_state.result = None

# ---- Input form ----
with st.form("job_agent_form"):
    col1, col2 = st.columns(2)
    with col1:
        st.markdown("**Your Resume**")
        resume_file = st.file_uploader("Upload PDF or DOCX", type=["pdf", "docx"], key="resume_upload")
        resume_text_manual = st.text_area(
            "...or paste resume text directly",
            height=200,
            placeholder="Paste your resume text here (optional if you uploaded a file above)...",
        )
    with col2:
        jd_text = st.text_area(
            "Job Description (paste as plain text)",
            height=300,
            placeholder="Paste the job description here...",
        )

    qa_question = st.text_input(
        "Optional: an application question you want help answering",
        placeholder="e.g. Why do you want to work here?",
    )

    submitted = st.form_submit_button("Analyze & Generate", type="primary")

if submitted:
    # Prefer uploaded file text; fall back to manually pasted text
    resume_text = ""
    if resume_file is not None:
        resume_text = extract_text_from_file(resume_file)
    if not resume_text.strip():
        resume_text = resume_text_manual

    if not resume_text.strip() or not jd_text.strip():
        st.error("Please upload or paste your resume, and paste the job description.")
    else:
        with st.spinner("Running multi-agent analysis... this can take 30-60s on a free-tier API key."):
            try:
                result = run_pipeline(
                    resume_text=resume_text,
                    jd_text=jd_text,
                    qa_question=qa_question.strip() if qa_question.strip() else None,
                )
                st.session_state.result = result
            except Exception as e:
                st.error(f"Pipeline failed: {e}")
                st.session_state.result = None

# ---- Results display ----
result = st.session_state.result
if result:
    if result.get("errors"):
        with st.expander("⚠️ Some agents reported errors (click to view)"):
            for err in result["errors"]:
                st.warning(err)

    st.divider()

    # --- Skill Gap Section ---
    st.subheader(f"📊 Skill Match: {result.get('match_percentage', 0)}%")
    st.progress(min(result.get("match_percentage", 0) / 100, 1.0))

    col1, col2 = st.columns(2)
    with col1:
        st.markdown("**✅ Matched Skills**")
        matched = result.get("matched_skills", [])
        if matched:
            st.write(", ".join(matched))
        else:
            st.write("_None identified_")
    with col2:
        st.markdown("**❌ Missing Skills**")
        missing = result.get("missing_skills", [])
        if missing:
            st.write(", ".join(missing))
        else:
            st.write("_None — great match!_")

    if result.get("gap_summary"):
        st.info(result["gap_summary"])

    st.divider()

    # --- Cover Letter Section ---
    st.subheader("✉️ Tailored Cover Letter")
    cover_letter = result.get("cover_letter", "")
    if cover_letter:
        st.text_area("Cover letter (editable, copy from here)", value=cover_letter, height=250)
    else:
        st.warning("Cover letter generation failed — check the errors panel above.")

    st.divider()

    # --- Resume Bullets Section ---
    st.subheader("📝 ATS-Optimized Resume Bullets")
    bullets = result.get("suggested_bullets", [])
    if bullets:
        for b in bullets:
            st.markdown(f"- {b}")
        st.caption("Bullets marked [ADD METRIC] need a real number from you — don't leave placeholders in your final resume.")
    else:
        st.warning("Bullet generation failed — check the errors panel above.")

    # --- Q&A Section (only if a question was asked) ---
    if result.get("qa_answer"):
        st.divider()
        st.subheader("💬 Application Question — Suggested Answer")
        st.text_area("Answer (editable)", value=result["qa_answer"], height=200)